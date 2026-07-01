from docx import Document
from pathlib import Path
import os

def document_maker(parameters: dict, response=None, player=None, session_memory=None) -> str:
    """
    Creates and saves a Word document (.docx) or Text file (.txt).
    """
    filename = parameters.get("filename", "INDRA_document").strip()
    format_type = parameters.get("format", "docx").lower()
    title = parameters.get("title", "")
    content = parameters.get("content", "")
    
    if not content:
        return "Cannot create a document without content."
        
    desktop = Path.home() / "Desktop"
    
    if format_type == "docx":
        if not filename.endswith(".docx"):
            filename += ".docx"
        file_path = desktop / filename
        
        try:
            doc = Document()
            if title:
                doc.add_heading(title, 0)

            paragraphs = content.split('\n\n')
            for p in paragraphs:
                if p.strip():
                    doc.add_paragraph(p.strip())
                    
            doc.save(str(file_path))
            return f"Successfully created Word document at: {file_path}"
        except Exception as e:
            return f"Failed to create .docx document: {e}"
            
    elif format_type == "txt":
        if not filename.endswith(".txt"):
            filename += ".txt"
        file_path = desktop / filename
        
        try:
            with open(file_path, "w", encoding="utf-8") as f:
                if title:
                    f.write(f"{title}\n{'='*len(title)}\n\n")
                f.write(content)
            return f"Successfully created Text document at: {file_path}"
        except Exception as e:
            return f"Failed to create .txt document: {e}"
            
    else:
        return f"Unknown document format: {format_type}. Supported: docx, txt."
